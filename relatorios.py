"""
relatorios.py — Geração dos entregáveis: relatório HTML e checklist XLSX.
"""
import io
from datetime import datetime


AUTORES = "Kevin Dias Quintian &nbsp;·&nbsp; Renata Gomes Rocha &nbsp;·&nbsp; Sergio Rosenboim &nbsp;·&nbsp; Viviane Nishizaki Suzuke &nbsp;·&nbsp; William Felipe dos Santos Moura"
RODAPE_TXT = "Kevin Dias Quintian · Renata Gomes Rocha · Sergio Rosenboim · Viviane Nishizaki Suzuke · William Felipe dos Santos Moura"

def gerar_relatorio_html(resultado: dict, modelo_nome: str) -> str:
    """Relatório HTML autocontido — identidade visual access Ai (mesmo CSS do app)."""
    # Import local: mantém relatorios.py importável sem a interface carregada
    from html import escape
    from ui_style import CSS, LOGO_SIMBOLO_SVG, status_badge
    from ui_componentes import (conformidade_cards_html, analise_geral_html, kpi, linha_kpis,
                                SIMBOLO_STATUS, SIGNIFICADO_STATUS, ROTULO_STATUS_PLURAL)
    from verificacoes import classificar_status

    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    resumo = resultado.get("resumo", {})
    itens = resultado.get("resultados", [])

    cores_txt = {"Conforme": "c-green", "Parcial": "c-purple", "Não Conforme": "c-red",
                 "Indeterminado": "c-amber", "N/A": "c-muted"}
    chaves = {"Conforme": "conformes", "Parcial": "parciais", "Não Conforme": "nao_conformes",
              "Indeterminado": "indeterminados", "N/A": "na"}
    cards_status = linha_kpis([
        kpi(f'<span class="{cores_txt[s_]}" style="font-size:1.1rem;margin-right:6px">{SIMBOLO_STATUS[s_]}</span>'
            f'{resumo.get(chaves[s_], 0)}', ROTULO_STATUS_PLURAL[s_], SIGNIFICADO_STATUS[s_].capitalize())
        for s_ in ["Conforme", "Parcial", "Não Conforme", "Indeterminado", "N/A"]])

    rows = ""
    for it in itens:
        gid = it.get("globalid", "") or ""
        gid_cell = f'<span class="globalid">{escape(gid)}</span>' if gid and gid != "—" else '<span class="td-muted">—</span>'
        stt = it.get("status", "N/A")
        rec = escape(it.get("recomendacao", "") or "") or "—"
        precisa_acao = classificar_status(stt) in ("Não Conforme", "Parcial", "Indeterminado")
        rec_cell = f'<span class="td-rec">{rec}</span>' if precisa_acao else f'<span class="td-muted">{rec}</span>'
        conferir = (' <span class="pill" title="Avaliação baseada no nome/tipo do elemento, não em medição direta">'
                    'conferir</span>') if it.get("requer_confirmacao_humana") else ""
        rows += f"""
        <tr>
          <td class="td-item">{escape(str(it.get('item_nbr', '—')))}{conferir}</td>
          <td class="td-muted">{escape(str(it.get('categoria', '—')))}</td>
          <td class="td-strong">{escape(str(it.get('elemento', '—')))}</td>
          <td>{status_badge(stt)}</td>
          <td class="td-mono" style="white-space:normal">{escape(str(it.get('valor_encontrado', '—')))}</td>
          <td class="td-mono" style="white-space:normal">{escape(str(it.get('valor_exigido', '—')))}</td>
          <td>{gid_cell}</td>
          <td class="td-mono td-muted">{escape(str(it.get('tipo_ifc', '—')))}</td>
          <td>{rec_cell}</td>
        </tr>"""

    return f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Relatório NBR 9050 — {escape(modelo_nome)}</title>
{CSS}
<style>
  body {{ margin: 0; background: var(--canvas); color: var(--deep); font-family: var(--font); }}
  .pagina {{ max-width: 1180px; margin: 0 auto; padding: 2.5rem 2rem 1rem 2rem; }}
  .rel-topo {{ display: flex; justify-content: space-between; align-items: flex-start; gap: 1.5rem;
              padding-bottom: 1.25rem; border-bottom: 1px solid var(--linha); margin-bottom: 1.25rem; flex-wrap: wrap; }}
  .rel-titulo {{ font-family: var(--display); font-size: 1.7rem; font-weight: 600; letter-spacing: -0.02em; margin-top: 1rem; }}
  .meta {{ display: flex; flex-wrap: wrap; gap: 0.5rem 1.5rem; font-size: 0.8rem; color: var(--muted); margin-bottom: 1.5rem; }}
  .meta strong {{ color: var(--deep); font-weight: 600; }}
  .nota-tabela {{ font-size: 0.78rem; color: var(--muted); margin: -0.25rem 0 0.75rem 0; line-height: 1.6; }}
  .rodape {{ max-width: 1180px; margin: 2.5rem auto 0 auto; padding: 1.25rem 2rem 2rem 2rem; border-top: 1px solid var(--linha);
            display: flex; justify-content: space-between; gap: 1.5rem; flex-wrap: wrap; font-size: 0.72rem; color: var(--muted); line-height: 1.6; }}
  .rodape strong {{ color: var(--texto-2); }}
  .rodape img {{ height: 18px; opacity: 0.8; display: block; margin-bottom: 0.35rem; margin-left: auto; }}
  @media print {{
    body {{ background: #fff; }}
    .metric-card, .conf-card, .analise, .result-table {{ box-shadow: none; }}
    thead {{ display: table-header-group; }}
    tr {{ page-break-inside: avoid; }}
  }}
</style>
</head>
<body>
<div class="pagina">

  <div class="rel-topo">
    <div>
      <div class="brand" style="border:none;padding:0;margin:0">
        <div class="brand-row">
          <span class="brand-word">access</span>
          <span class="brand-mark" style="height:24px;width:27px">{LOGO_SIMBOLO_SVG}</span>
        </div>
        <div class="brand-tagline">NBR 9050 · Accessibility Checker</div>
      </div>
      <div class="rel-titulo">Relatório de verificação de acessibilidade</div>
    </div>
    <div class="hero-sub" style="text-align:right;margin-top:0.4rem">Auditoria BIM<br>ABNT NBR 9050:2020</div>
  </div>

  <div class="meta">
    <span>Modelo: <strong>{escape(modelo_nome)}</strong></span>
    <span>Schema IFC: <strong>{escape(str(resultado.get('schema_ifc', '—')))}</strong></span>
    <span>Norma: <strong>ABNT NBR 9050:2020</strong></span>
    <span>Emitido em: <strong>{now}</strong></span>
  </div>

  <div class="section-title">Conformidade do modelo</div>
  {conformidade_cards_html(resumo)}

  <div class="section-title">Itens da norma por resultado <span class="section-sub">{resumo.get('total', len(itens))} itens avaliados</span></div>
  {cards_status}

  <div class="section-title">Análise geral</div>
  {analise_geral_html(resultado)}

  <div class="section-title">Resultados detalhados</div>
  <div class="nota-tabela">
    <strong>GlobalId</strong> é o identificador único do elemento no IFC (o "CPF" do elemento): use-o no Revit,
    Navisworks ou BIMcollab para localizar o elemento no modelo.
    <span class="pill">conferir</span> indica avaliação feita pelo nome do elemento, e não por medição — recomenda-se confirmação humana.
    <strong>Parcial</strong> significa que parte dos elementos atende e parte não (ver "Encontrado").
  </div>
  <div style="overflow-x:auto">
  <table class="result-table">
    <thead>
      <tr>
        <th>Item NBR</th><th>Categoria</th><th>Elemento</th><th>Status</th>
        <th>Encontrado</th><th>Exigido</th><th>GlobalId</th><th>Tipo IFC</th><th>Recomendação</th>
      </tr>
    </thead>
    <tbody>{rows}</tbody>
  </table>
  </div>

</div>

<div class="rodape">
  <div>
    <strong>TFM · Grupo 1</strong><br>
    {AUTORES}<br>
    Gerado automaticamente — itens qualitativos e indeterminados exigem verificação manual complementar.
  </div>
  <div style="text-align:right">
    <img src="https://www.e-zigurat.com/images/logo.svg" alt="Zigurat Institute of Technology" />
    Master Internacional em IA para Arquitetura e Construção<br>
    Zigurat Institute of Technology · {now}
  </div>
</div>

</body>
</html>"""


def gerar_excel(resultado: dict, modelo_nome: str) -> bytes:
    """Generate XLSX report with openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    except ImportError:
        return b""

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Checklist NBR 9050"

    # Colors
    H_FILL  = PatternFill("solid", fgColor="0A1628")
    SUB_FILL = PatternFill("solid", fgColor="111827")
    CONF_F  = PatternFill("solid", fgColor="064E3B")
    PARC_F  = PatternFill("solid", fgColor="3B1064")
    NAO_F   = PatternFill("solid", fgColor="4C0519")
    INDET_F = PatternFill("solid", fgColor="4B2B06")
    NA_F    = PatternFill("solid", fgColor="1E293B")

    thin = Side(style="thin", color="1F2D45")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    # Header
    ws.merge_cells("A1:I1")
    ws["A1"] = f"♿ RELATÓRIO DE ACESSIBILIDADE NBR 9050:2020 — {modelo_nome}"
    ws["A1"].font = Font(bold=True, color="00D4AA", size=12, name="Calibri")
    ws["A1"].fill = H_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:I2")
    ws["A2"] = f"Emitido em: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Schema IFC: {resultado.get('schema_ifc','—')}"
    ws["A2"].font = Font(color="64748B", size=9, italic=True, name="Calibri")
    ws["A2"].fill = SUB_FILL
    ws["A2"].alignment = Alignment(horizontal="center")

    # Column headers
    headers = ["Item NBR", "Categoria", "Elemento", "Status", "Valor Encontrado",
               "Valor Exigido", "GlobalId (Revit/IFC)", "Tipo IFC", "Recomendação"]
    widths   = [12, 16, 32, 16, 20, 20, 32, 18, 45]

    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=3, column=i, value=h)
        cell.font = Font(bold=True, color="94A3B8", size=9, name="Calibri")
        cell.fill = SUB_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[chr(64+i)].width = w
    ws.row_dimensions[3].height = 20

    # Data rows
    status_fills = {
        "conforme": CONF_F, "parcial": PARC_F, "não conforme": NAO_F, "nao conforme": NAO_F,
        "indeterminado": INDET_F, "n/a": NA_F
    }
    status_colors = {
        "conforme": "10B981", "parcial": "C084FC", "não conforme": "F87171", "nao conforme": "F87171",
        "indeterminado": "FCD34D", "n/a": "64748B"
    }

    for r, it in enumerate(resultado.get("resultados", []), start=4):
        st_key = it.get("status","").lower()
        fill = status_fills.get(st_key, NA_F)

        values = [
            it.get("item_nbr",""),
            it.get("categoria",""),
            it.get("elemento",""),
            it.get("status",""),
            it.get("valor_encontrado",""),
            it.get("valor_exigido",""),
            it.get("globalid","—"),
            it.get("tipo_ifc",""),
            it.get("recomendacao",""),
        ]
        for c, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.fill = fill
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(color="E2E8F0", size=9, name="Calibri")
            if c == 4:
                color = status_colors.get(st_key, "E2E8F0")
                cell.font = Font(color=color, bold=True, size=9, name="Calibri")
            if c == 7:
                cell.font = Font(color="60A5FA", size=9, name="Calibri Mono")
        ws.row_dimensions[r].height = 36

    # Summary sheet
    ws2 = wb.create_sheet("Resumo")
    resumo = resultado.get("resumo", {})
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 20
    ws2.append(["Métrica", "Valor"])
    ws2.append(["Modelo", resultado.get("modelo","—")])
    ws2.append(["Schema IFC", resultado.get("schema_ifc","—")])
    ws2.append(["Data Auditoria", resultado.get("data_auditoria", datetime.now().strftime("%d/%m/%Y"))])
    ws2.append(["Total de Itens", resumo.get("total", 0)])
    ws2.append(["✅ Conformes", resumo.get("conformes", 0)])
    ws2.append(["▲ Parciais", resumo.get("parciais", 0)])
    ws2.append(["❌ Não Conformes", resumo.get("nao_conformes", 0)])
    ws2.append(["⚠️ Indeterminados", resumo.get("indeterminados", 0)])
    ws2.append(["— N/A", resumo.get("na", 0)])
    ws2.append(["% Conformidade", resumo.get("percentual_conformidade","—")])
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row):
        for cell in row:
            cell.font = Font(name="Calibri", size=10)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
